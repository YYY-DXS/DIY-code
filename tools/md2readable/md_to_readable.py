#!/usr/bin/env python3
"""把项目里的 Markdown 转成两种易读格式：HTML（浏览器直接看）和 DOCX（Word 编辑）。

用法:
    python md_to_readable.py            # 转换 projects/ 下所有 md
    python md_to_readable.py --html     # 只生成 HTML
    python md_to_readable.py --docx     # 只生成 DOCX
"""

from __future__ import annotations

import argparse
import html as html_lib
import re
import sys
from datetime import datetime
from pathlib import Path

# ---------------------------------------------------------------- 解析 Markdown


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)" )


def parse_inline(text: str) -> list[tuple[str, str]]:
    """拆成 [('text', 内容), ('bold', 内容), ('code', 内容), ('link', 文本|地址)]"""
    parts: list[tuple[str, str]] = []
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            parts.append(("bold", token[2:-2]))
        elif token.startswith("`") and token.endswith("`"):
            parts.append(("code", token[1:-1]))
        elif token.startswith("[") and "](" in token:
            label, _, url = token[1:-1].partition("](")
            parts.append(("link", f"{label}\u0000{url}"))
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            parts.append(("italic", token[1:-1]))
        else:
            parts.append(("text", token))
    return parts


def parse_markdown(text: str) -> dict:
    lines = text.splitlines()
    doc: dict = {"title": "", "meta": [], "blocks": []}
    table: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table
        if table:
            doc["blocks"].append({"type": "table", "rows": table})
            table = []

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            table.append(cells)
            continue
        flush_table()

        if not stripped:
            continue
        if stripped == "---":
            continue
        if stripped.startswith("# "):
            doc["title"] = stripped[2:].strip()
        elif stripped.startswith("## "):
            doc["blocks"].append({"type": "h2", "text": stripped[3:].strip()})
        elif stripped.startswith("### "):
            doc["blocks"].append({"type": "h3", "text": stripped[4:].strip()})
        elif stripped.startswith("> "):
            doc["blocks"].append({"type": "quote", "text": stripped[2:].strip()})
        elif re.match(r"^- \[[ xX]\] ", stripped):
            done = stripped[3] in "xX"
            doc["blocks"].append(
                {"type": "check", "text": stripped[6:].strip(), "done": done}
            )
        elif stripped.startswith("- "):
            doc["blocks"].append({"type": "bullet", "text": stripped[2:].strip()})
        elif re.match(r"^\d+\. ", stripped):
            doc["blocks"].append(
                {"type": "number", "text": re.sub(r"^\d+\. ", "", stripped)}
            )
        elif not doc["title"] and not doc["blocks"]:
            doc["meta"].append(stripped)
        else:
            doc["blocks"].append({"type": "p", "text": stripped})
    flush_table()
    return doc


# ---------------------------------------------------------------- HTML


HTML_CSS = """
:root { color-scheme: light; }
body { margin: 0; background: #f6f7f9; color: #1f2328;
  font-family: "Microsoft YaHei", "PingFang SC", "Noto Sans CJK SC", system-ui, sans-serif;
  font-size: 16px; line-height: 1.85; }
main { max-width: 820px; margin: 0 auto; padding: 48px 28px 80px;
  background: #fff; box-shadow: 0 1px 24px rgba(0,0,0,.06); }
h1 { font-size: 30px; line-height: 1.35; margin: 0 0 8px; font-weight: 700; }
.meta { color: #6b7280; font-size: 14px; margin-bottom: 28px; }
h2 { font-size: 21px; margin: 40px 0 12px; padding-bottom: 6px;
  border-bottom: 1px solid #e5e7eb; font-weight: 700; }
h3 { font-size: 17px; margin: 26px 0 8px; font-weight: 700; }
p { margin: 12px 0; }
ul, ol { margin: 12px 0; padding-left: 26px; }
li { margin: 6px 0; }
li.check { list-style: none; margin-left: -22px; }
li.check::before { content: "☐ "; color: #6b7280; }
li.check.done { color: #9ca3af; text-decoration: line-through; }
li.check.done::before { content: "☑ "; }
blockquote { margin: 18px 0; padding: 2px 0 2px 16px; border-left: 3px solid #d1d5db;
  color: #4b5563; }
code { background: #f3f4f6; padding: 1px 5px; border-radius: 4px;
  font-family: Consolas, "Courier New", monospace; font-size: 14px; }
table { width: 100%; border-collapse: collapse; margin: 18px 0; font-size: 15px; }
th, td { border: 1px solid #d9d9d9; padding: 9px 12px; text-align: left; vertical-align: middle; }
th { background: #eef2f6; font-weight: 700; }
tbody tr:nth-child(even) { background: #fafbfc; }
a { color: #1a56db; text-decoration: none; }
a:hover { text-decoration: underline; }
.footer { margin-top: 56px; padding-top: 16px; border-top: 1px solid #e5e7eb;
  color: #9ca3af; font-size: 13px; }
@media print { body { background: #fff; } main { box-shadow: none; max-width: none; padding: 0; } }
"""


def link_target(url: str) -> str:
    if url.startswith("http"):
        return url
    if url.endswith(".md"):
        return url[:-3] + ".html"
    return url


def inline_html(text: str) -> str:
    out = []
    for kind, value in parse_inline(text):
        if kind == "bold":
            out.append(f"<strong>{html_lib.escape(value)}</strong>")
        elif kind == "code":
            out.append(f"<code>{html_lib.escape(value)}</code>")
        elif kind == "italic":
            out.append(f"<em>{html_lib.escape(value)}</em>")
        elif kind == "link":
            label, _, url = value.partition("\u0000")
            out.append(
                f'<a href="{html_lib.escape(link_target(url))}">{html_lib.escape(label)}</a>'
            )
        else:
            out.append(html_lib.escape(value))
    return "".join(out)


def to_html(doc: dict, source_name: str) -> str:
    body: list[str] = [f"<h1>{html_lib.escape(doc['title'] or source_name)}</h1>"]
    if doc["meta"]:
        body.append(f'<p class="meta">{html_lib.escape(" ".join(doc["meta"]))}</p>')
    open_list: str | None = None

    def close_list() -> None:
        nonlocal open_list
        if open_list:
            body.append(f"</{open_list}>")
            open_list = None

    for block in doc["blocks"]:
        kind = block["type"]
        if kind in ("bullet", "check", "number"):
            want = "ol" if kind == "number" else "ul"
            if open_list != want:
                close_list()
                body.append(f"<{want}>")
                open_list = want
            if kind == "check":
                cls = "check done" if block["done"] else "check"
                body.append(f'<li class="{cls}">{inline_html(block["text"])}</li>')
            else:
                body.append(f"<li>{inline_html(block['text'])}</li>")
            continue
        close_list()
        if kind == "h2":
            body.append(f"<h2>{inline_html(block['text'])}</h2>")
        elif kind == "h3":
            body.append(f"<h3>{inline_html(block['text'])}</h3>")
        elif kind == "quote":
            body.append(f"<blockquote>{inline_html(block['text'])}</blockquote>")
        elif kind == "table":
            rows = block["rows"]
            body.append("<table>")
            for index, row in enumerate(rows):
                tag = "th" if index == 0 else "td"
                cells = "".join(f"<{tag}>{inline_html(c)}</{tag}>" for c in row)
                body.append(f"<tr>{cells}</tr>")
            body.append("</table>")
        else:
            body.append(f"<p>{inline_html(block['text'])}</p>")
    close_list()
    body.append(
        f'<p class="footer">由 Markdown 自动转换 · {datetime.now().strftime("%Y-%m-%d")}</p>'
    )
    return (
        "<!doctype html>\n<html lang=\"zh-CN\">\n<head>\n<meta charset=\"utf-8\">\n"
        f"<title>{html_lib.escape(doc['title'] or source_name)}</title>\n"
        f"<style>{HTML_CSS}</style>\n</head>\n<body>\n<main>\n"
        + "\n".join(body)
        + "\n</main>\n</body>\n</html>\n"
    )


# ---------------------------------------------------------------- DOCX


def set_cjk_font(run, name: str = "DengXian") -> None:
    from docx.oxml.ns import qn

    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)


def add_runs(paragraph, text: str) -> None:
    for kind, value in parse_inline(text):
        if kind == "link":
            label, _, _url = value.partition("\u0000")
            run = paragraph.add_run(label)
        else:
            run = paragraph.add_run(value)
            if kind == "bold":
                run.bold = True
            elif kind == "italic":
                run.italic = True
            elif kind == "code":
                run.font.name = "Consolas"
        set_cjk_font(run)


def style_table(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(3)
                para.paragraph_format.space_after = Pt(3)
                for run in para.runs:
                    set_cjk_font(run)
                    if row_index == 0:
                        run.bold = True
            shade = "EEF2F6" if row_index == 0 else (
                "FAFBFC" if row_index % 2 == 0 else "FFFFFF"
            )
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), shade)
            tc_pr.append(shd)


def to_docx(doc: dict, source_name: str, out_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "DengXian"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    title = document.add_paragraph(style="Title")
    run = title.add_run(doc["title"] or source_name)
    set_cjk_font(run, "DengXian")
    run.font.size = Pt(24)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    if doc["meta"]:
        meta = document.add_paragraph()
        meta_run = meta.add_run(" ".join(doc["meta"]))
        set_cjk_font(meta_run)
        meta_run.font.size = Pt(10)
        meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    for block in doc["blocks"]:
        kind = block["type"]
        if kind in ("h2", "h3"):
            para = document.add_paragraph(style="Heading 1" if kind == "h2" else "Heading 2")
            run = para.add_run(block["text"])
            set_cjk_font(run)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.size = Pt(16 if kind == "h2" else 13)
            run.bold = True
        elif kind == "quote":
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Pt(18)
            run = para.add_run(block["text"])
            set_cjk_font(run)
            run.italic = True
            run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        elif kind == "table":
            rows = block["rows"]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = document.add_table(rows=0, cols=width)
            for row in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row[:width]):
                    paragraph = cells[index].paragraphs[0]
                    add_runs(paragraph, value)
            style_table(table)
            document.add_paragraph()
        elif kind == "check":
            para = document.add_paragraph(style="List Bullet")
            run = para.add_run(("☑ " if block["done"] else "☐ ") + block["text"])
            set_cjk_font(run)
            if block["done"]:
                run.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
        elif kind == "bullet":
            para = document.add_paragraph(style="List Bullet")
            add_runs(para, block["text"])
        elif kind == "number":
            para = document.add_paragraph(style="List Number")
            add_runs(para, block["text"])
        else:
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(para, block["text"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))


# ---------------------------------------------------------------- 主流程


def main() -> int:
    ap = argparse.ArgumentParser(description="Markdown 转易读格式")
    ap.add_argument("--root", default=None, help="要处理的目录（默认仓库的 projects/）")
    ap.add_argument("--html", action="store_true", help="只生成 HTML")
    ap.add_argument("--docx", action="store_true", help="只生成 DOCX")
    args = ap.parse_args()

    repo = Path(__file__).resolve().parents[2]
    root = Path(args.root) if args.root else repo / "projects"
    if not root.exists():
        print(f"找不到目录：{root}")
        return 1

    make_html = args.html or not args.docx
    make_docx = args.docx or not args.html

    files = sorted(root.rglob("*.md"))
    if not files:
        print("没有找到 markdown 文件")
        return 1

    done = 0
    for md in files:
        doc = parse_markdown(md.read_text(encoding="utf-8"))
        base = md.with_suffix("")
        if make_html:
            Path(f"{base}.html").write_text(
                to_html(doc, md.stem), encoding="utf-8"
            )
        if make_docx:
            to_docx(doc, md.stem, Path(f"{base}.docx"))
        done += 1
        print(f"  转换：{md.relative_to(root)}")
    print(f"\n共处理 {done} 个文件（HTML={make_html} DOCX={make_docx}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
